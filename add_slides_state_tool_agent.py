"""
Appends State, Tool, and Agent listing slides to Document_Assistant.pptx.
Run: python add_slides_state_tool_agent.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PPTX_PATH = "Document_Assistant.pptx"

DARK_BG = RGBColor(0x1E, 0x1E, 0x2E)
ACCENT  = RGBColor(0x74, 0xC7, 0xEC)
GREEN   = RGBColor(0xA6, 0xE3, 0xA1)
YELLOW  = RGBColor(0xF9, 0xE2, 0xAF)
PURPLE  = RGBColor(0xCB, 0xA6, 0xF7)
RED     = RGBColor(0xF3, 0x8B, 0xA8)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT   = RGBColor(0xCD, 0xD6, 0xF4)
ORANGE  = RGBColor(0xFA, 0xB3, 0x87)


# ── Helpers ────────────────────────────────────────────────────────────────────

def add_slide(prs):
    layout = prs.slide_layouts[0]
    slide  = prs.slides.add_slide(layout)
    for ph in slide.placeholders:
        ph._element.getparent().remove(ph._element)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG
    return slide


def title_box(slide, text):
    txb = slide.shapes.add_textbox(Inches(0.4), Inches(0.18), Inches(9.2), Inches(0.65))
    p   = txb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = ACCENT


def divider_line(slide, y=Inches(0.88)):
    ln = slide.shapes.add_shape(1, Inches(0.4), y, Inches(9.2), Inches(0.03))
    ln.fill.solid()
    ln.fill.fore_color.rgb = ACCENT
    ln.line.fill.background()


def text_box(slide, text, left, top, width, height,
             size=Pt(12), colour=WHITE, bold=False, mono=False, align=PP_ALIGN.LEFT):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf  = txb.text_frame
    tf.word_wrap = True
    first = True
    for line in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size  = size
        run.font.bold  = bold
        run.font.color.rgb = colour
        if mono:
            run.font.name = "Courier New"


def code_box(slide, code, left, top, width, height, size=Pt(10.5)):
    rect = slide.shapes.add_shape(1, left, top, width, height)
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(0x11, 0x11, 0x22)
    rect.line.fill.background()
    txb = slide.shapes.add_textbox(
        left + Inches(0.12), top + Inches(0.1),
        width - Inches(0.24), height - Inches(0.15))
    tf = txb.text_frame
    tf.word_wrap = True
    first = True
    for line in code.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = line
        run.font.size = size
        run.font.name = "Courier New"
        run.font.color.rgb = GREEN


def badge(slide, label, left, top, w=Inches(0.28), h=Inches(0.28), colour=ACCENT):
    rect = slide.shapes.add_shape(1, left, top, w, h)
    rect.fill.solid()
    rect.fill.fore_color.rgb = colour
    rect.line.fill.background()
    txb = slide.shapes.add_textbox(left, top, w, h)
    p   = txb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = DARK_BG


def row(slide, y, col1, col2, col3, col4,
        c1=YELLOW, c2=ACCENT, c3=WHITE, c4=LIGHT, h=Inches(0.38)):
    widths = [Inches(1.7), Inches(1.3), Inches(2.8), Inches(3.2)]
    lefts  = [Inches(0.35), Inches(2.1), Inches(3.45), Inches(6.3)]
    for text, left, width, colour in zip(
            [col1, col2, col3, col4], lefts, widths, [c1, c2, c3, c4]):
        txb = slide.shapes.add_textbox(left, y, width, h)
        tf  = txb.text_frame
        tf.word_wrap = True
        p   = tf.paragraphs[0]
        run = p.add_run()
        run.text  = text
        run.font.size  = Pt(11)
        run.font.color.rgb = colour
        run.font.bold = (colour in (YELLOW, ACCENT))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE A — AgentState full listing
# ══════════════════════════════════════════════════════════════════════════════
def slide_agent_state(prs):
    slide = add_slide(prs)
    title_box(slide, "AgentState — Full Field Listing (agent.py)")
    divider_line(slide)

    # Header row
    hdr_y = Inches(0.95)
    for text, left, width, colour in zip(
        ["Field", "Type", "Reducer", "Purpose"],
        [Inches(0.35), Inches(2.1), Inches(3.45), Inches(6.3)],
        [Inches(1.7),  Inches(1.3), Inches(2.8),  Inches(3.2)],
        [ACCENT, ACCENT, ACCENT, ACCENT]
    ):
        txb = slide.shapes.add_textbox(left, hdr_y, width, Inches(0.32))
        p   = txb.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = text
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = colour

    fields = [
        ("user_input",           "str",              "—",                  "Current user message text"),
        ("messages",             "List[BaseMessage]","add_messages",        "Full conversation history (append-only)"),
        ("intent",               "UserIntent | None","—",                  "Classified intent from classify_intent node"),
        ("next_step",            "str",              "—",                  "Target node name for routing"),
        ("conversation_summary", "str",              "—",                  "LLM-generated rolling summary"),
        ("active_documents",     "List[str]",        "—",                  "Document IDs referenced across turns"),
        ("current_response",     "Pydantic | None",  "—",                  "Structured output from active agent"),
        ("tools_used",           "List[str]",        "—",                  "Tool names called this turn"),
        ("session_id",           "str",              "—",                  "LangGraph thread_id for checkpointing"),
        ("user_id",              "str",              "—",                  "User identifier"),
        ("actions_taken",        "List[str]",        "operator.add",       "Node names run this turn (accumulated)"),
    ]

    start_y = Inches(1.32)
    row_h   = Inches(0.43)
    for i, (f, t, r, p) in enumerate(fields):
        y = start_y + i * row_h
        # alternating background
        if i % 2 == 0:
            bg = slide.shapes.add_shape(1, Inches(0.3), y, Inches(9.4), row_h - Inches(0.03))
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(0x28, 0x28, 0x40)
            bg.line.fill.background()
        row(slide, y + Inches(0.05), f, t, r, p,
            c1=YELLOW, c2=ORANGE, c3=GREEN, c4=WHITE, h=row_h - Inches(0.08))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE B — Reducer explanation
# ══════════════════════════════════════════════════════════════════════════════
def slide_reducers(prs):
    slide = add_slide(prs)
    title_box(slide, "AgentState — Reducers Explained")
    divider_line(slide)

    code = """\
# Without a reducer — last write wins (default)
state["next_step"] = "qa_agent"          # classify_intent sets this
state["next_step"] = "summarization"     # update_memory overwrites it  ← lost!


# add_messages reducer — appends, never overwrites
messages: Annotated[List[BaseMessage], add_messages]

# Turn 1 result : [HumanMessage("Q1 revenue?"), AIMessage("$4.2M")]
# Turn 2 result : [HumanMessage("Q2 revenue?"), AIMessage("$5.1M")]
# State after 2 turns: ALL 4 messages preserved  ✓


# operator.add reducer — list concatenation across nodes
actions_taken: Annotated[List[str], operator.add]

# classify_intent returns : ["classify_intent"]
# qa_agent returns        : ["qa_agent"]
# update_memory returns   : ["update_memory"]
# Final state             : ["classify_intent", "qa_agent", "update_memory"]  ✓


# InMemorySaver — checkpoints full state after every node
workflow.compile(checkpointer=InMemorySaver())
# Same thread_id → state restored on next .invoke()  →  true multi-turn memory"""

    code_box(slide, code, Inches(0.35), Inches(1.0), Inches(9.3), Inches(5.7))


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE C — Tools listing
# ══════════════════════════════════════════════════════════════════════════════
def slide_tools_overview(prs):
    slide = add_slide(prs)
    title_box(slide, "Tools — calculator & document_reader (tools.py)")
    divider_line(slide)

    # calculator block
    text_box(slide, "🧮  calculator", Inches(0.35), Inches(1.0),
             Inches(9.2), Inches(0.35), size=Pt(14), colour=YELLOW, bold=True)

    calc = """\
@tool
def calculator(expression: str) -> str:
    # 1. Strip commas:        "4,200,000 + 5,100,000"  →  "4200000 + 5100000"
    # 2. Regex whitelist:     only  digits  + - * / ( ) . %  allowed
    # 3. Sandboxed eval():    eval(expr, {"__builtins__": {}}, {})
    # 4. Format result:       int-valued floats → "9,300,000"
    # 5. Log to session file: status = OK | INVALID_INPUT | ZERO_DIVISION | ERROR
    return "Result: 9,300,000  (expression: 4200000 + 5100000)" """
    code_box(slide, calc, Inches(0.35), Inches(1.4), Inches(9.3), Inches(2.0))

    # document_reader block
    text_box(slide, "📄  document_reader", Inches(0.35), Inches(3.5),
             Inches(9.2), Inches(0.35), size=Pt(14), colour=YELLOW, bold=True)

    dr = """\
@tool
def document_reader(query: str) -> str:
    # Path 1  LIST        query = "list"         → returns all doc IDs + titles
    # Path 2  EXACT_MATCH query = "doc_financial_q2"  → returns full content
    # Path 3  KEYWORD     query = "revenue growth"    → keyword scored search (top 3)
    # Path 4  FALLBACK    no match found         → returns first 2 docs in store
    # Path 5  NO_MATCH    store is empty         → "No documents found"
    # All paths log to logs/<session_id>.log with status label + full output"""
    code_box(slide, dr, Inches(0.35), Inches(3.9), Inches(9.3), Inches(1.85))

    text_box(slide,
             "Both tools are bound to the LLM via llm.bind_tools([calculator, document_reader])\n"
             "and invoked inside _run_tool_loop() when the LLM returns a tool_calls response.",
             Inches(0.35), Inches(5.82), Inches(9.3), Inches(0.55),
             size=Pt(11), colour=LIGHT)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE D — Tool flow diagram
# ══════════════════════════════════════════════════════════════════════════════
def slide_tool_loop(prs):
    slide = add_slide(prs)
    title_box(slide, "Tool Calling Loop — _run_tool_loop()")
    divider_line(slide)

    flow = """\
def _run_tool_loop(llm_with_tools, tool_map, messages, max_iter=5):
    tools_used      = []
    context_snippets = []

    for _ in range(max_iter):          # max 5 iterations (8 for calculation_agent)

        response = llm_with_tools.invoke(messages)
        messages.append(response)

        if not response.tool_calls:    # LLM finished — no more tool calls
            break

        for tc in response.tool_calls:
            name   = tc["name"]        # e.g. "document_reader"
            result = tool_map[name].invoke(tc["args"])   # runs the tool
            tools_used.append(name)
            context_snippets.append(str(result))
            messages.append(ToolMessage(content=result, tool_call_id=tc["id"]))

    return messages, tools_used, context_snippets


# If context_snippets is still empty after the loop:
if not context_snippets:
    context_snippets = _direct_retrieve(state["user_input"])
    # → searches retriever directly and logs a FALLBACK entry"""
    code_box(slide, flow, Inches(0.35), Inches(1.0), Inches(9.3), Inches(4.3))

    text_box(slide,
             "Each tool result becomes a ToolMessage appended to the message chain "
             "so the LLM sees it in the next iteration.\n"
             "After the loop, context_snippets feeds the Phase 2 structured-output call.",
             Inches(0.35), Inches(5.4), Inches(9.3), Inches(0.65),
             size=Pt(11), colour=LIGHT)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE E — All 5 agents/nodes listing
# ══════════════════════════════════════════════════════════════════════════════
def slide_agents_listing(prs):
    slide = add_slide(prs)
    title_box(slide, "All 5 Agent Nodes — Listing & Role (agent.py)")
    divider_line(slide)

    nodes = [
        ("1", "classify_intent",      ACCENT,
         "WHEN: Entry point — every message\n"
         "DOES: Calls llm.with_structured_output(UserIntent) with the enriched\n"
         "      intent classification prompt + last 6 messages as history.\n"
         "OUT:  intent_type, confidence, reasoning → sets next_step for routing"),

        ("2", "qa_agent",             GREEN,
         "WHEN: intent_type == 'qa'\n"
         "DOES: Tool loop with document_reader → Phase 2 with AnswerResponse.\n"
         "      Injects _available_docs_hint() so LLM knows all doc IDs.\n"
         "OUT:  AnswerResponse(question, answer, sources, confidence)"),

        ("3", "summarization_agent",  YELLOW,
         "WHEN: intent_type == 'summarization'\n"
         "DOES: Tool loop retrieves document → Phase 2 with SummarizationResponse.\n"
         "      Extracts 3–7 key points from document content.\n"
         "OUT:  SummarizationResponse(summary, key_points, sources)"),

        ("4", "calculation_agent",    ORANGE,
         "WHEN: intent_type == 'calculation'\n"
         "DOES: Tool loop (max 8 iter) — calls document_reader then calculator.\n"
         "      Phase 2 with CalculationResponse including the expression used.\n"
         "OUT:  CalculationResponse(question, result, expression, sources)"),

        ("5", "update_memory",        PURPLE,
         "WHEN: After any specialist agent — always the last node before END\n"
         "DOES: LLM writes 1–2 sentence conversation summary. Tracks active_documents\n"
         "      by reading sources from current_response. Sets next_step = END.\n"
         "OUT:  conversation_summary, active_documents updated in state"),
    ]

    start_y = Inches(1.0)
    block_h = Inches(1.02)
    for num, name, colour, desc in nodes:
        y = start_y + (int(num) - 1) * block_h

        # Background block
        bg = slide.shapes.add_shape(1, Inches(0.3), y, Inches(9.4), block_h - Inches(0.06))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(0x22, 0x22, 0x38)
        bg.line.color.rgb = colour
        bg.line.width = Pt(1)

        # Number badge
        badge(slide, num, Inches(0.38), y + Inches(0.3), colour=colour)

        # Node name
        text_box(slide, name, Inches(0.75), y + Inches(0.25),
                 Inches(2.0), Inches(0.45), size=Pt(13), colour=colour, bold=True, mono=True)

        # Description
        text_box(slide, desc, Inches(2.85), y + Inches(0.08),
                 Inches(6.75), Inches(block_h - Inches(0.18)),
                 size=Pt(10.5), colour=WHITE)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE F — Graph routing diagram (text-based)
# ══════════════════════════════════════════════════════════════════════════════
def slide_graph_routing(prs):
    slide = add_slide(prs)
    title_box(slide, "LangGraph Routing — Edges & Conditional Flow")
    divider_line(slide)

    diagram = """\
                        ┌─────────────────────────┐
                        │   User sends a message   │
                        └────────────┬────────────┘
                                     │
                                     ▼
                          ┌──────────────────────┐
                          │   classify_intent     │  ← Entry Point
                          │   (UserIntent schema) │
                          └──────────┬───────────┘
                                     │
               ┌─────────────────────┼──────────────────────┐
               │ intent="qa"         │ intent="summarization" │ intent="calculation"
               ▼                     ▼                        ▼
       ┌──────────────┐   ┌─────────────────────┐   ┌──────────────────────┐
       │   qa_agent   │   │ summarization_agent  │   │  calculation_agent   │
       │  AnswerResp  │   │ SummarizationResp    │   │  CalculationResp     │
       └──────┬───────┘   └──────────┬──────────┘   └──────────┬───────────┘
              │                      │                           │
              └──────────────────────┼───────────────────────────┘
                                     │
                                     ▼
                           ┌──────────────────┐
                           │  update_memory   │  ← Always runs last
                           │  (summary + docs)│
                           └────────┬─────────┘
                                    │
                                    ▼
                                   END


Graph compiled with: workflow.compile(checkpointer=InMemorySaver())
Routing function:    _route_to_agent(state) → returns state["next_step"]"""

    code_box(slide, diagram, Inches(0.35), Inches(1.0), Inches(9.3), Inches(5.7), size=Pt(10))


# ══════════════════════════════════════════════════════════════════════════════
# Main
# ══════════════════════════════════════════════════════════════════════════════
def main():
    prs = Presentation(PPTX_PATH)
    before = len(prs.slides)

    slide_agent_state(prs)
    slide_reducers(prs)
    slide_tools_overview(prs)
    slide_tool_loop(prs)
    slide_agents_listing(prs)
    slide_graph_routing(prs)

    prs.save(PPTX_PATH)
    print(f"Added {len(prs.slides) - before} slides. Total: {len(prs.slides)}")


if __name__ == "__main__":
    main()
